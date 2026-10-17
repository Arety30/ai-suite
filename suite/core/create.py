"""Create workspace — AI-generated documents, slide decks, and spreadsheets,
exported to real local files.

Content is produced by the local LLM (Ollama); export uses local libraries
(python-docx, fpdf2, python-pptx, openpyxl) — no cloud, no keys. Files land in
data/exports and are fetched back via /v1/create/download/{name}.

Structured artifacts (slides, sheets) are generated with a coder model in JSON
mode because small chat models don't emit clean structure; prose docs use the
default chat chain.
"""
from __future__ import annotations

import csv
import json
import re

from suite.config import settings
from suite.core.documents import _first_json_block  # reuse robust JSON extraction
from suite.providers.base import registry

EXPORTS = settings.data_dir / "exports"
STRUCT_MODEL = settings.raw.get("create", {}).get("model", "qwen2.5-coder:3b")


def _slug(title: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", title.strip().lower()).strip("-")
    return s or "untitled"


def _path(title: str, ext: str):
    EXPORTS.mkdir(parents=True, exist_ok=True)
    # ordinal suffix avoids clobbering same-titled exports
    base = _slug(title)
    n, path = 0, EXPORTS / f"{base}.{ext}"
    while path.exists():
        n += 1
        path = EXPORTS / f"{base}-{n}.{ext}"
    return path


def _material(doc_ids: list[int] | None) -> str:
    """Fetch the text of vault sources so generation can be grounded in them."""
    if not doc_ids:
        return ""
    from suite.core import memory
    with memory.connect() as db:
        marks = ",".join("?" * len(doc_ids))
        rows = db.execute(
            f"SELECT title, text FROM document WHERE id IN ({marks})", doc_ids).fetchall()
    return "\n\n".join(f"[{r['title']}]\n{r['text']}" for r in rows)


def _with_sources(prompt: str, material: str) -> str:
    if not material:
        return prompt
    return (f"Base this strictly on the following sources:\n\n{material[:6000]}\n\n"
            f"Task: {prompt}")


async def _prose(prompt: str, material: str = "") -> str:
    msgs = [{"role": "system", "content": "Write clear, well-structured content in "
             "Markdown (use # headings and - bullets). Output only the document."},
            {"role": "user", "content": _with_sources(prompt, material)}]
    reply = await registry.call("llm", "chat", msgs)
    return reply.text


async def _gen_json(prompt: str, shape: str, material: str = ""):
    """Return the raw parsed JSON (dict or list) the model produced."""
    msgs = [{"role": "system", "content": shape},
            {"role": "user", "content": _with_sources(prompt, material)}]
    reply = await registry.call("llm", "chat", msgs, model=STRUCT_MODEL, format="json")
    block = _first_json_block(reply.text) or reply.text
    try:
        return json.loads(block)
    except json.JSONDecodeError:
        return {}


# ---- Docs -------------------------------------------------------------------

def _md_to_docx(title: str, md: str):
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=0)
    for line in md.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith(("- ", "* ")):
            doc.add_paragraph(s[2:], style="List Bullet")
        else:
            doc.add_paragraph(s)
    path = _path(title, "docx")
    doc.save(path)
    return path


def _md_to_pdf(title: str, md: str):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    # core fonts are latin-1; sanitize so unicode from the LLM never crashes export
    def w(text, size, style=""):
        pdf.set_font("Helvetica", style, size)
        pdf.multi_cell(0, 8, text.encode("latin-1", "replace").decode("latin-1"))
    w(title, 18, "B")
    pdf.ln(2)
    for line in md.splitlines():
        s = line.strip()
        if not s:
            pdf.ln(2); continue
        if s.startswith("# "): w(s[2:], 15, "B")
        elif s.startswith("## "): w(s[3:], 13, "B")
        elif s.startswith("### "): w(s[4:], 12, "B")
        elif s.startswith(("- ", "* ")): w(chr(149) + " " + s[2:], 11)
        else: w(s, 11)
    path = _path(title, "pdf")
    pdf.output(str(path))
    return path


async def make_doc(user_id: int, title: str, prompt: str, fmt: str = "docx",
                   doc_ids: list[int] | None = None) -> dict:
    md = await _prose(prompt, _material(doc_ids))
    if fmt == "pdf":
        path = _md_to_pdf(title, md)
    elif fmt == "md":
        path = _path(title, "md"); path.write_text(f"# {title}\n\n{md}")
    else:
        path = _md_to_docx(title, md)
    return {"file": path.name, "title": title, "kind": "doc", "format": fmt}


# ---- Slides -----------------------------------------------------------------

_SLIDE_SHAPE = ('Return a JSON object {"slides":[{"title": "...", "bullets": '
                '["...","..."]}]} with 5-8 slides. No prose outside the JSON.')


async def make_slides(user_id: int, title: str, prompt: str,
                      doc_ids: list[int] | None = None) -> dict:
    data = await _gen_json(prompt, _SLIDE_SHAPE, _material(doc_ids))
    slides = data.get("slides", []) if isinstance(data, dict) else (data or [])
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    # title slide
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = title
    for sl in slides:
        layout = prs.slide_layouts[1]  # title + content
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = str(sl.get("title", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = sl.get("bullets") or []
        for i, b in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(b)
            para.font.size = Pt(18)
    path = _path(title, "pptx")
    prs.save(path)
    return {"file": path.name, "title": title, "kind": "slides", "slides": len(slides)}


# ---- Sheets -----------------------------------------------------------------

_SHEET_SHAPE = ('Return a JSON object {"headers":["..."],"rows":[["..."]]} — a '
                'table with a header row and data rows. No prose outside the JSON.')


async def make_sheet(user_id: int, title: str, prompt: str, fmt: str = "xlsx",
                     doc_ids: list[int] | None = None) -> dict:
    data = await _gen_json(prompt, _SHEET_SHAPE, _material(doc_ids))
    headers, rows = _coerce_table(data)
    if fmt == "csv":
        path = _path(title, "csv")
        with open(path, "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(headers)
            w.writerows(rows)
    else:
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = title[:31] or "Sheet"
        ws.append(headers)
        for r in rows:
            ws.append([str(c) for c in r])
        path = _path(title, "xlsx")
        wb.save(path)
    return {"file": path.name, "title": title, "kind": "sheet",
            "rows": len(rows), "cols": len(headers)}


def _coerce_table(data) -> tuple[list, list]:
    """Normalize the model's JSON into (headers, rows). Handles {headers,rows},
    a list of row-objects, or a bare list of rows."""
    if isinstance(data, dict):
        if "headers" in data or "rows" in data:
            return data.get("headers", []), data.get("rows", [])
        # single {col: val} object -> one row
        return list(data.keys()), [list(data.values())]
    items = data or []
    if items and isinstance(items[0], dict):
        headers = list(items[0].keys())
        return headers, [[row.get(h, "") for h in headers] for row in items]
    if items and isinstance(items[0], list):
        return [f"col{i+1}" for i in range(len(items[0]))], items
    return ["value"], [[i] for i in items]


def list_exports() -> list[dict]:
    EXPORTS.mkdir(parents=True, exist_ok=True)
    return [{"file": p.name, "size": p.stat().st_size}
            for p in sorted(EXPORTS.iterdir(), key=lambda x: x.stat().st_mtime, reverse=True)
            if p.is_file()]
